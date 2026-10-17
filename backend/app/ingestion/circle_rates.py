"""Circle rate ingestion pipeline for the UP / Lucknow pilot.

Downloads official circle rates (guideline values) from the state portal,
normalizes them into the circle_rates schema, and upserts them into
PostgreSQL. Never fabricates rates: a missing or misconfigured source
raises an error instead of inserting synthetic values.

F-04: this pipeline must not silently drop rows. Every received row is accounted
for, and any row that cannot be normalized is emitted as a reason-coded
``QuarantineRecord`` next to an ``IngestionRunAccounting`` record, matching the
canonical ingestion pipeline (hard-rule 13: rejected rows are logged with a reason,
never silently fixed).
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import ipaddress
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from uuid import UUID, uuid4

import httpx
from bs4 import BeautifulSoup
from sqlalchemy.dialects.postgresql import insert as pg_insert

from app.core.config import get_settings
from app.core.logging import get_logger
from app.database.connection import AsyncSessionLocal
from app.database.models.circle_rate import CircleRate
from app.ingestion.provider import BUILD_VERSION, compute_sha256
from app.ingestion.quarantine import (
    IngestionRunAccounting,
    QuarantineReason,
    QuarantineRecord,
)
from app.ingestion.registry import generate_source_id
from app.ingestion.schemas import CircleRateRecord

logger = get_logger(__name__)

DEFAULT_QUARANTINE_DIR = Path(__file__).resolve().parents[3] / "data" / "quarantine"

# F-17: canonical default property type, aligned with the reference-rate pipeline and the
# circle-rates lookup API default (was the ad-hoc "residential").
DEFAULT_PROPERTY_TYPE = "residential_flat"

# F-08: file types accepted for local (offline) ingestion.
ALLOWED_INTAKE_SUFFIXES: frozenset[str] = frozenset({".csv", ".xlsx", ".xls", ".pdf"})


def _is_blocked_ip_literal(host: str) -> bool:
    """Whether a hostname is an IP literal in an internal/reserved range (F-08)."""
    try:
        address = ipaddress.ip_address(host)
    except ValueError:
        return False
    return (
        address.is_private
        or address.is_loopback
        or address.is_link_local
        or address.is_multicast
        or address.is_reserved
        or address.is_unspecified
    )


def _host_is_allowed(host: str, allowed_hosts: tuple[str, ...]) -> bool:
    """Exact or subdomain match against the approved host allowlist."""
    hostname = host.lower().rstrip(".")
    return any(
        hostname == allowed or hostname.endswith(f".{allowed}")
        for allowed in allowed_hosts
        if allowed
    )


def validate_ingestion_url(url: str, allowed_hosts: tuple[str, ...]) -> str:
    """Validate an outbound ingestion URL against the approved host policy (F-08).

    Fails closed: HTTPS-only, IP literals in internal/reserved ranges are rejected, and the
    host must be on the configured allowlist. No host is assumed to be "official" because
    provenance is not yet verified (see F-03), so an empty allowlist refuses to fetch.
    """
    if not url or not url.strip():
        raise IngestionError(
            "CIRCLE_RATE_SOURCE_URL is not configured. Set it to the official portal URL "
            "before running ingestion."
        )
    parsed = urlparse(url.strip())
    if parsed.scheme != "https":
        raise IngestionError(
            f"Refusing ingestion URL with scheme {parsed.scheme!r}; HTTPS is required (F-08)."
        )
    host = (parsed.hostname or "").lower()
    if not host:
        raise IngestionError("Ingestion URL has no host (F-08).")
    if _is_blocked_ip_literal(host):
        raise IngestionError(
            f"Refusing ingestion URL pointing at an internal/reserved address ({host!r}) (F-08)."
        )
    if not allowed_hosts:
        raise IngestionError(
            "No approved ingestion hosts configured (CIRCLE_RATE_ALLOWED_HOSTS is empty); "
            "refusing to fetch (F-08)."
        )
    if not _host_is_allowed(host, allowed_hosts):
        raise IngestionError(
            f"Host {host!r} is not in the approved ingestion host allowlist (F-08)."
        )
    return url.strip()


def resolve_intake_path(path: str | Path, intake_dir: str | Path) -> Path:
    """Resolve a local ingestion path inside the approved intake directory (F-08).

    Prevents path traversal and reading arbitrary files outside the approved directory.
    """
    candidate = Path(path).expanduser()
    if not candidate.is_absolute():
        candidate = Path.cwd() / candidate
    resolved = candidate.resolve()
    approved = Path(intake_dir).expanduser().resolve()
    if not resolved.is_relative_to(approved):
        raise IngestionError(
            f"Refusing to read {resolved}: outside the approved intake directory {approved} (F-08)."
        )
    if resolved.suffix.lower() not in ALLOWED_INTAKE_SUFFIXES:
        raise IngestionError(f"Unsupported file type: {resolved.suffix} (F-08).")
    return resolved

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "state": ("state", "state_name", "state code"),
    "district": ("district", "dist", "district_name"),
    "zone_tehsil": ("tehsil", "zone", "zone_tehsil", "tehsil_name", "taluka"),
    "locality": ("locality", "area", "village", "mouza", "locality_name"),
    "property_type": ("property_type", "type of property", "category", "usage"),
    "rate_per_sqm": (
        "rate per sqm",
        "rate/sqm",
        "circle rate",
        "circle rate per sqm",
        "guideline value",
        "value per sqm",
    ),
    "effective_year": ("effective year", "year", "from year", "notification year"),
}

SOURCE_COLUMNS = (
    "state",
    "district",
    "zone_tehsil",
    "locality",
    "property_type",
    "rate_per_sqm",
    "effective_year",
    "source_url",
    "fetched_at",
)


@dataclass(frozen=True)
class CircleRateRejection:
    """A received circle-rate row that could not be normalized, with a reason code."""

    raw_ref: str
    reason_code: QuarantineReason
    reason_detail: str


@dataclass
class ParseOutcome:
    """Result of parsing a snapshot: accepted rows plus reason-coded rejections.

    ``rows_received == len(rows) + len(rejections)`` always holds, so no received
    row can be silently lost.
    """

    rows: list[dict[str, Any]] = field(default_factory=list)
    rejections: list[CircleRateRejection] = field(default_factory=list)
    rows_received: int = 0


@dataclass
class CircleRateIngestionResult:
    """Outcome of one circle-rate ingestion run (parse + validate; DB write is separate)."""

    records: list[CircleRateRecord]
    accounting: IngestionRunAccounting
    quarantine: list[QuarantineRecord]
    quarantine_path: Path | None = None
    rows_upserted: int = 0


class IngestionError(Exception):
    """Raised when circle rate ingestion cannot complete safely."""


class CircleRateIngestor:
    """Fetch, normalize, and persist official circle rates."""

    def __init__(
        self,
        source_url: str | None = None,
        state: str = "UP",
        district: str = "Lucknow",
        city: str = "Lucknow",
        quarantine_dir: str | Path | None = None,
        intake_dir: str | Path | None = None,
        allowed_hosts: tuple[str, ...] | list[str] | None = None,
    ) -> None:
        settings = get_settings()
        self.source_url = source_url or settings.CIRCLE_RATE_SOURCE_URL
        self.state = state.strip().upper()
        self.district = district.strip()
        self.city = city.strip()
        self.quarantine_dir = (
            Path(quarantine_dir) if quarantine_dir is not None else DEFAULT_QUARANTINE_DIR
        )
        # F-08: approved intake directory and outbound host allowlist.
        self.intake_dir = (
            Path(intake_dir) if intake_dir is not None else Path(settings.CIRCLE_RATE_INTAKE_DIR)
        )
        if allowed_hosts is None:
            configured = settings.CIRCLE_RATE_ALLOWED_HOSTS
            if isinstance(configured, str):  # pragma: no cover - validator returns a list
                configured = [configured]
            self.allowed_hosts: tuple[str, ...] = tuple(
                str(host).strip().lower() for host in configured if str(host).strip()
            )
        else:
            self.allowed_hosts = tuple(str(host).strip().lower() for host in allowed_hosts if str(host).strip())

    @property
    def source_id(self) -> UUID:
        """Deterministic source UUID used for accounting and quarantine provenance."""
        return generate_source_id(
            self.state, self.source_url or f"local://circle-rates/{self.state}"
        )

    async def fetch_raw(self) -> str:
        """Download the circle rate lookup page from the configured source.

        F-08: the URL is validated against the approved host policy and redirects are **not**
        followed, so a compromised/redirecting host cannot pivot to an internal address.
        """
        url = validate_ingestion_url(self.source_url, self.allowed_hosts)
        async with httpx.AsyncClient(timeout=30.0, follow_redirects=False) as client:
            response = await client.get(url)
            if response.is_redirect:
                location = response.headers.get("location", "")
                raise IngestionError(
                    f"Refusing to follow redirect from {url} to {location!r} (F-08 host policy)."
                )
            response.raise_for_status()
            return response.text

    def parse_html(self, raw_html: str) -> ParseOutcome:
        """Parse circle rate tables from a portal HTML page.

        IGRS-style portals publish rates in HTML tables keyed by district /
        tehsil / locality. Header names vary between portal releases, so
        verify the alias map below against the live page before scheduling
        monthly runs.

        F-04: rows that are malformed or cannot be normalized are recorded as
        rejections instead of being skipped silently.
        """
        soup = BeautifulSoup(raw_html, "html.parser")
        tables = soup.find_all("table")
        outcome = ParseOutcome()

        for table in tables:
            headers = [th.get_text(" ", strip=True).lower() for th in table.find_all("th")]
            if not headers or not _looks_like_rate_table(headers):
                continue
            for row in table.find_all("tr")[1:]:
                cells = [td.get_text(" ", strip=True) for td in row.find_all("td")]
                if not cells:
                    continue  # structural spacer row with no data cells
                outcome.rows_received += 1
                raw_ref = _cells_ref(cells)
                if len(cells) < max(len(headers), 1):
                    outcome.rejections.append(
                        CircleRateRejection(
                            raw_ref=raw_ref,
                            reason_code=QuarantineReason.MISSING_REQUIRED_FIELD,
                            reason_detail=(
                                f"Row has {len(cells)} cells but table declares {len(headers)} columns"
                            ),
                        )
                    )
                    continue
                record = dict(zip(headers, cells, strict=False))
                _collect(outcome, record, self.state, self.district, self.source_url, raw_ref)

        if not outcome.rows and not outcome.rejections:
            raise IngestionError("No circle rate tables recognized in the portal page.")

        return outcome

    def parse_file(self, path: str) -> ParseOutcome:
        """Parse circle rates from a local CSV, Excel, or PDF file (offline testing).

        F-08: the path must resolve inside the approved intake directory.
        """
        resolved = resolve_intake_path(path, self.intake_dir)
        suffix = resolved.suffix.lower()
        path_str = str(resolved)
        if suffix == ".csv":
            return self._parse_csv(path_str)
        if suffix in (".xlsx", ".xls"):
            return self._parse_excel(path_str)
        if suffix == ".pdf":
            return self._parse_pdf(path_str)
        raise IngestionError(f"Unsupported file type: {suffix}")

    def _parse_csv(self, path: str) -> ParseOutcome:
        outcome = ParseOutcome()
        source_url = self.source_url or "file://" + path
        with open(path, newline="", encoding="utf-8") as handle:
            reader = csv.DictReader(handle)
            for row in reader:
                outcome.rows_received += 1
                record = {str(k).lower().strip(): v for k, v in row.items() if k is not None}
                if None in row:
                    outcome.rejections.append(
                        CircleRateRejection(
                            raw_ref=_row_ref(record),
                            reason_code=QuarantineReason.MISSING_REQUIRED_FIELD,
                            reason_detail="Row has more fields than the header row",
                        )
                    )
                    continue
                _collect(outcome, record, self.state, self.district, source_url, _row_ref(record))
        return outcome

    def _parse_excel(self, path: str) -> ParseOutcome:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        outcome = ParseOutcome()
        source_url = self.source_url or "file://" + path
        table_found = False
        for sheet in workbook.worksheets:
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                continue
            headers = [str(h).lower().strip() if h is not None else "" for h in values[0]]
            if not _looks_like_rate_table(headers):
                continue
            table_found = True
            for raw_row in values[1:]:
                if all(value is None or not str(value).strip() for value in raw_row):
                    continue  # blank spacer row
                outcome.rows_received += 1
                record = dict(zip(headers, raw_row, strict=False))
                _collect(outcome, record, self.state, self.district, source_url, _row_ref(record))
        if not table_found:
            raise IngestionError(f"No circle rate tables recognized in workbook: {path}")
        return outcome

    def _parse_pdf(self, path: str) -> ParseOutcome:
        import pdfplumber

        outcome = ParseOutcome()
        source_url = self.source_url or "file://" + path
        table_found = False
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                table = page.extract_table()
                if not table or len(table) < 2:
                    continue
                headers = [str(h).lower().strip() if h is not None else "" for h in table[0]]
                if not _looks_like_rate_table(headers):
                    continue
                table_found = True
                for raw_row in table[1:]:
                    if all(value is None or not str(value).strip() for value in raw_row):
                        continue  # blank spacer row
                    outcome.rows_received += 1
                    record = dict(zip(headers, raw_row, strict=False))
                    _collect(
                        outcome, record, self.state, self.district, source_url, _row_ref(record)
                    )
        if not table_found:
            raise IngestionError(f"No circle rate tables recognized in PDF: {path}")
        return outcome

    def validate_records(
        self, rows: list[dict[str, Any]]
    ) -> tuple[list[CircleRateRecord], list[CircleRateRejection]]:
        """Coerce normalized rows into ``CircleRateRecord`` objects.

        F-04: a schema failure is returned as a rejection (never swallowed).
        """
        validated: list[CircleRateRecord] = []
        rejections: list[CircleRateRejection] = []
        now = datetime.now(UTC)
        for row in rows:
            try:
                record = CircleRateRecord(
                    state=row["state"],
                    district=row["district"],
                    zone_tehsil=row["zone_tehsil"],
                    locality=row["locality"],
                    property_type=row["property_type"],
                    rate_per_sqm=row["rate_per_sqm"],
                    effective_year=row["effective_year"],
                    source_url=row["source_url"],
                    fetched_at=row.get("fetched_at") or now,
                )
            except Exception as exc:  # noqa: BLE001 - converted to a reason-coded rejection
                rejections.append(
                    CircleRateRejection(
                        raw_ref=_row_ref(row),
                        reason_code=QuarantineReason.MISSING_REQUIRED_FIELD,
                        reason_detail=f"Schema validation failed: {exc}",
                    )
                )
                continue
            validated.append(record)
        return validated, rejections

    def parse_records(self, rows: list[dict[str, Any]]) -> list[CircleRateRecord]:
        """Backwards-compatible wrapper returning only validated records.

        Prefer :meth:`ingest_html` / :meth:`ingest_file`, which also return the
        reason-coded quarantine records and run accounting.
        """
        validated, _ = self.validate_records(rows)
        return validated

    def _finalize(
        self,
        outcome: ParseOutcome,
        snapshot_hash: str,
        *,
        run_id: UUID | None = None,
        persist_quarantine: bool = True,
    ) -> CircleRateIngestionResult:
        """Convert a parse outcome into records + accounting + quarantine log."""
        records, record_rejections = self.validate_records(outcome.rows)
        rejections = [*outcome.rejections, *record_rejections]
        resolved_run_id = run_id or uuid4()
        source_id = self.source_id

        quarantine = [
            QuarantineRecord(
                run_id=resolved_run_id,
                source_id=source_id,
                source_record_id=rejection.raw_ref,
                reason_code=rejection.reason_code,
                reason_detail=rejection.reason_detail,
                raw_record_ref=rejection.raw_ref,
            )
            for rejection in rejections
        ]

        accounting = IngestionRunAccounting(
            run_id=resolved_run_id,
            source_id=source_id,
            started_at=datetime.now(UTC),
            finished_at=datetime.now(UTC),
            build_version=BUILD_VERSION,
            source_snapshot_hash=snapshot_hash,
            rows_received=outcome.rows_received,
            rows_accepted=len(records),
            rows_quarantined=len(quarantine),
            rows_duplicate=0,
            rows_rejected=0,
        )

        quarantine_path: Path | None = None
        if quarantine and persist_quarantine:
            self.quarantine_dir.mkdir(parents=True, exist_ok=True)
            quarantine_path = self.quarantine_dir / f"quarantine_{resolved_run_id}.jsonl"
            with open(quarantine_path, mode="w", encoding="utf-8") as handle:
                for record in quarantine:
                    handle.write(record.model_dump_json() + "\n")

        return CircleRateIngestionResult(
            records=records,
            accounting=accounting,
            quarantine=quarantine,
            quarantine_path=quarantine_path,
        )

    def ingest_html(
        self,
        raw_html: str,
        *,
        run_id: UUID | None = None,
        persist_quarantine: bool = True,
    ) -> CircleRateIngestionResult:
        """Parse and validate an HTML snapshot with full row accounting (F-04)."""
        outcome = self.parse_html(raw_html)
        return self._finalize(
            outcome,
            compute_sha256(raw_html),
            run_id=run_id,
            persist_quarantine=persist_quarantine,
        )

    def ingest_file(
        self,
        path: str | Path,
        *,
        run_id: UUID | None = None,
        persist_quarantine: bool = True,
    ) -> CircleRateIngestionResult:
        """Parse and validate a local file snapshot with full row accounting (F-04/F-08)."""
        resolved_path = resolve_intake_path(path, self.intake_dir)
        outcome = self.parse_file(str(resolved_path))
        return self._finalize(
            outcome,
            compute_sha256(resolved_path.read_bytes()),
            run_id=run_id,
            persist_quarantine=persist_quarantine,
        )

    async def upsert_records(self, session, records: list[CircleRateRecord]) -> int:
        """Insert or update circle rates keyed by the natural unique constraint."""
        if not records:
            return 0

        count = 0
        for record in records:
            insert_stmt = pg_insert(CircleRate).values(
                state=record.state,
                district=record.district,
                zone_tehsil=record.zone_tehsil,
                locality=record.locality,
                property_type=record.property_type,
                rate_per_sqm=record.rate_per_sqm,
                effective_year=record.effective_year,
                source_url=record.source_url,
                fetched_at=record.fetched_at,
            )
            # `excluded` is an attribute of the *statement*, not of the `pg_insert` callable.
            stmt = insert_stmt.on_conflict_do_update(
                constraint="uq_circle_rate_state_district_zone_locality_type_year",
                set_={
                    "rate_per_sqm": insert_stmt.excluded.rate_per_sqm,
                    "source_url": insert_stmt.excluded.source_url,
                    "fetched_at": insert_stmt.excluded.fetched_at,
                },
            )
            await session.execute(stmt)
            count += 1
        await session.commit()
        return count

    async def run(self) -> CircleRateIngestionResult:
        """Execute the full ingestion pipeline and return the accounted run result (F-04)."""
        raw_html = await self.fetch_raw()
        result = self.ingest_html(raw_html)
        accounting = result.accounting
        if accounting.rows_quarantined:
            logger.warning(
                "Quarantined %d circle rate row(s) from %s (log: %s)",
                accounting.rows_quarantined,
                self.source_url,
                result.quarantine_path,
            )

        async with AsyncSessionLocal() as session:
            result.rows_upserted = await self.upsert_records(session, result.records)
        logger.info(
            "Upserted %d/%d circle rates for state=%s district=%s",
            result.rows_upserted,
            accounting.rows_accepted,
            self.state,
            self.district,
        )
        return result


def _looks_like_rate_table(headers: list[str]) -> bool:
    joined = " ".join(headers)
    has_locality = any(alias in joined for alias in HEADER_ALIASES["locality"])
    has_rate = any(alias in joined for alias in HEADER_ALIASES["rate_per_sqm"])
    return has_locality and has_rate


def _row_ref(row: dict[str, Any]) -> str:
    """Best-effort human-readable reference to a raw row for quarantine audit."""
    for key in ("locality", "area", "village", "mouza", "locality_name", "tehsil", "zone_tehsil"):
        value = str(row.get(key, "")).strip()
        if value:
            return value[:120]
    joined = " | ".join(str(value) for value in row.values() if str(value).strip())
    return joined[:120] or "unknown-row"


def _cells_ref(cells: list[str]) -> str:
    """Reference to a raw HTML table row (list of cell strings)."""
    joined = " | ".join(str(cell) for cell in cells if str(cell).strip())
    return joined[:120] or "unknown-row"


def _collect(
    outcome: ParseOutcome,
    row: dict[str, Any],
    state: str,
    district: str,
    source_url: str,
    raw_ref: str,
) -> None:
    """Normalize one received row, recording a reason-coded rejection on failure (F-04)."""
    normalized, reason, detail = _normalize_row_with_reason(row, state, district, source_url)
    if normalized is None:
        outcome.rejections.append(
            CircleRateRejection(
                raw_ref=raw_ref,
                reason_code=reason or QuarantineReason.MISSING_REQUIRED_FIELD,
                reason_detail=detail or "Row could not be normalized",
            )
        )
        return
    outcome.rows.append(normalized)


def _normalize_row(
    row: dict[str, Any], state: str, district: str, source_url: str
) -> dict[str, Any] | None:
    """Normalize a raw row, returning the record or None (backwards-compatible helper)."""
    return _normalize_row_with_reason(row, state, district, source_url)[0]


def _normalize_row_with_reason(
    row: dict[str, Any], state: str, district: str, source_url: str
) -> tuple[dict[str, Any] | None, QuarantineReason | None, str]:
    """Normalize a raw row, returning ``(record, None, "")`` or ``(None, reason, detail)``."""
    mapped: dict[str, str] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        for header, value in row.items():
            if header in aliases and canonical not in mapped:
                mapped[canonical] = value
                break

    if "state" not in mapped and state:
        mapped["state"] = state
    if "district" not in mapped and district:
        mapped["district"] = district
    if "property_type" not in mapped:
        mapped["property_type"] = DEFAULT_PROPERTY_TYPE

    missing = [key for key in SOURCE_COLUMNS[:7] if key not in mapped or not str(mapped[key]).strip()]
    if missing:
        return None, QuarantineReason.MISSING_REQUIRED_FIELD, f"Missing required field(s): {', '.join(missing)}"

    rate = _clean_rate(mapped["rate_per_sqm"])
    if rate is None:
        return (
            None,
            QuarantineReason.MISSING_PRICE,
            f"Unparseable or non-positive circle rate: {mapped['rate_per_sqm']!r}",
        )
    year = _clean_year(mapped["effective_year"])
    if year is None:
        return (
            None,
            QuarantineReason.INVALID_DATE,
            f"Unparseable notification year: {mapped['effective_year']!r}",
        )

    return (
        {
            "state": state,
            "district": district,
            "zone_tehsil": str(mapped["zone_tehsil"]).strip(),
            "locality": str(mapped["locality"]).strip(),
            "property_type": str(mapped.get("property_type", DEFAULT_PROPERTY_TYPE)).strip()
            or DEFAULT_PROPERTY_TYPE,
            "rate_per_sqm": rate,
            "effective_year": year,
            "source_url": source_url,
            "fetched_at": datetime.now(UTC),
        },
        None,
        "",
    )


def _clean_rate(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).replace(",", "").replace("₹", "").replace("rs", "").strip()
    try:
        rate = float(text)
    except ValueError:
        return None
    return rate if rate > 0 else None


def _clean_year(value: Any) -> int | None:
    if value is None:
        return None
    text = str(value).strip()
    match = None
    import re

    match = re.search(r"(20\d{2})", text)
    if not match:
        return None
    return int(match.group(1))


def main() -> None:
    """CLI entry point for running circle rate ingestion."""
    parser = argparse.ArgumentParser(description="Ingest official circle rates into PostgreSQL.")
    parser.add_argument("--source", help="Portal URL (overrides CIRCLE_RATE_SOURCE_URL)")
    parser.add_argument("--state", default="UP", help="State code (default UP)")
    parser.add_argument("--district", default="Lucknow", help="District (default Lucknow)")
    parser.add_argument("--city", default="Lucknow", help="City (default Lucknow)")
    parser.add_argument("--file", help="Local CSV/Excel/PDF to ingest instead of fetching the portal")
    parser.add_argument("--dry-run", action="store_true", help="Parse and validate without writing")
    parser.add_argument(
        "--quarantine-dir",
        help="Directory for the quarantine JSONL log (default: data/quarantine)",
    )
    parser.add_argument(
        "--intake-dir",
        help="Approved directory for --file ingestion (default: CIRCLE_RATE_INTAKE_DIR)",
    )
    parser.add_argument(
        "--allow-host",
        action="append",
        default=None,
        help="Approved outbound host (repeatable; overrides CIRCLE_RATE_ALLOWED_HOSTS)",
    )
    args = parser.parse_args()

    async def _run() -> None:
        ingestor = CircleRateIngestor(
            source_url=args.source,
            state=args.state,
            district=args.district,
            city=args.city,
            quarantine_dir=args.quarantine_dir,
            intake_dir=args.intake_dir,
            allowed_hosts=args.allow_host,
        )
        if args.file:
            result = ingestor.ingest_file(args.file, persist_quarantine=not args.dry_run)
        else:
            raw = await ingestor.fetch_raw()
            result = ingestor.ingest_html(raw, persist_quarantine=not args.dry_run)

        accounting = result.accounting
        print(
            f"Row accounting: received={accounting.rows_received} "
            f"accepted={accounting.rows_accepted} "
            f"quarantined={accounting.rows_quarantined} "
            f"duplicate={accounting.rows_duplicate} "
            f"rejected={accounting.rows_rejected}"
        )
        if result.quarantine_path:
            print(f"Quarantine log: {result.quarantine_path}")

        if args.dry_run:
            for record in result.records[:10]:
                print(record)
            return

        async with AsyncSessionLocal() as session:
            count = await ingestor.upsert_records(session, result.records)
        print(f"Upserted {count} circle rates")

    asyncio.run(_run())


if __name__ == "__main__":
    main()
